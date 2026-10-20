"""Builds a professional, styled .xlsx export of a computed report -
overall summary, a detailed sheet with every item (target items bold and
highlighted, exactly like the original أصناف الأورانج report style), and
one sheet per pharmacy so branch managers can be sent just their page.

Quantity is always shown; القيمة / اجمالي القيمة (بالليرة السورية الجديدة)
are included whenever the upload has at least one known item price."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TITLE_FILL = PatternFill("solid", fgColor="1F3864")
SUBTITLE_FILL = PatternFill("solid", fgColor="EBF5FB")
NOTE_FILL = PatternFill("solid", fgColor="FFFDE7")
HEADER_FILL = PatternFill("solid", fgColor="2F5496")
CATEGORY_BAND_FILL = PatternFill("solid", fgColor="354F72")
TARGET_FILL = PatternFill("solid", fgColor="FFF2CC")
STRIPE_FILL = PatternFill("solid", fgColor="EBF5FB")
WARN_FILL = PatternFill("solid", fgColor="FDEBD0")
WHITE_FONT_BOLD = Font(color="FFFFFF", bold=True, size=13)
HEADER_FONT = Font(color="FFFFFF", bold=True)
CATEGORY_FONT = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CURRENCY_LABEL = "ل.س.ج"  # الليرة السورية الجديدة


def _detail_headers(has_prices):
    headers = ["#", "التصنيف", "الصنف", "الكمية", "% من التصنيف"]
    if has_prices:
        headers += [
            f"السعر ({CURRENCY_LABEL})",
            f"القيمة ({CURRENCY_LABEL})",
            "% القيمة من التصنيف",
        ]
    headers += ["مستهدف؟"]
    return headers


NO_FILL = PatternFill(fill_type=None)
PLAIN_FONT = Font()
CENTER = Alignment(horizontal="center")
CENTER_MID = Alignment(horizontal="center", vertical="center")


class SheetWriter:
    """A thin write-only wrapper that keeps the old sheet layout.

    The workbook used to be built cell by cell in openpyxl's normal mode,
    which keeps every Cell object in memory until the file is saved. A
    month with ~34,000 item rows, written once into the full-detail sheet
    and again into each pharmacy's own sheet, came to roughly 68,000
    styled rows and peaked near 370MB - more than a small instance has,
    so the export killed the worker and the browser got an HTTP 502.

    Write-only mode streams each row to disk and discards it, so peak
    memory no longer grows with the number of rows. The one feature it
    does not offer is merged cells, so the merged title/subtitle/category
    bands are written instead as a full-width coloured row: the text sits
    in the first cell and every remaining cell of that row carries the
    same fill, which looks the same on the page."""

    def __init__(self, workbook, title, num_cols, widths=None, freeze=None):
        self.ws = workbook.create_sheet(title)
        self.ws.sheet_view.rightToLeft = True
        self.num_cols = num_cols
        # Column widths and the frozen header must be set BEFORE the first
        # row is appended: a write-only sheet has already streamed its
        # <cols> and <sheetView> elements out by the time rows start, so
        # setting them afterwards is silently dropped and the file opens
        # with default widths and nothing frozen.
        if widths:
            for col, width in enumerate(widths, start=1):
                self.ws.column_dimensions[get_column_letter(col)].width = width
        if freeze:
            self.ws.freeze_panes = freeze

    def _cell(self, value, fill=None, font=None, alignment=None, border=None):
        cell = WriteOnlyCell(self.ws, value=value)
        if fill is not None:
            cell.fill = fill
        if font is not None:
            cell.font = font
        if alignment is not None:
            cell.alignment = alignment
        if border is not None:
            cell.border = border
        return cell

    def band(self, text, fill, font=None, height=None):
        """One full-width coloured row - the write-only stand-in for a
        merged cell."""
        row = [self._cell(text, fill=fill, font=font, alignment=CENTER_MID)]
        row += [self._cell(None, fill=fill) for _ in range(self.num_cols - 1)]
        self.ws.append(row)
        if height:
            self.ws.row_dimensions[self.ws._current_row].height = height

    def blank(self):
        self.ws.append([])

    def header(self, headers):
        self.ws.append([
            self._cell(h, fill=HEADER_FILL, font=HEADER_FONT, alignment=CENTER, border=BORDER)
            for h in headers
        ])

    def row(self, values, fill=None, font=None, center_cols=()):
        self.ws.append([
            self._cell(
                v, fill=fill, font=font, border=BORDER,
                alignment=CENTER if (col in center_cols) else None,
            )
            for col, v in enumerate(values, start=1)
        ])


def _sheet_title(sw, title, subtitle=None):
    sw.band(title, TITLE_FILL, WHITE_FONT_BOLD, height=26)
    if subtitle:
        sw.band(subtitle, SUBTITLE_FILL, BOLD)
    sw.band(
        "الصفوف الملوّنة الغامقة هي الأصناف المستهدفة (التي طلبت من الصيدلية بيعها)",
        NOTE_FILL,
    )
    sw.blank()


def _write_table(sw, rows_by_category, has_prices):
    headers = _detail_headers(has_prices)
    sw.header(headers)
    numeric_cols = {4, 5} if not has_prices else {4, 5, 6, 7, 8}

    counter = 1
    for cat_name, cat_total_qty, cat_total_value, items in rows_by_category:
        band_text = f"  {cat_name}   |   الكمية: {cat_total_qty:g} عبوة"
        if has_prices:
            band_text += f"   |   إجمالي القيمة: {cat_total_value:,.0f} {CURRENCY_LABEL}"
        sw.band(band_text, CATEGORY_BAND_FILL, CATEGORY_FONT)
        for i, item in enumerate(items):
            pct = (item.quantity / cat_total_qty * 100) if cat_total_qty else 0
            fill = TARGET_FILL if item.is_target else (STRIPE_FILL if i % 2 else NO_FILL)
            font = BOLD if item.is_target else PLAIN_FONT
            values = [counter, cat_name, item.item_name, round(item.quantity, 2), round(pct, 1)]
            if has_prices:
                value_pct = (item.value / cat_total_value * 100) if cat_total_value else 0
                values += [
                    round(item.unit_price, 2) if item.has_price else "—",
                    round(item.value, 0) if item.has_price else "—",
                    round(value_pct, 1) if item.has_price else "—",
                ]
            values += ["نعم" if item.is_target else ""]
            sw.row(values, fill=fill, font=font, center_cols=numeric_cols)
            counter += 1


def _detail_widths(has_prices):
    widths = [6, 30, 42, 12, 13]
    if has_prices:
        widths += [14, 16, 16]
    return widths + [10]


#: Every sheet opens with the same four-row banner (title, subtitle, note,
#: blank), so the header is row 5 and the data starts at row 6.
FREEZE_AT = "A6"


def _rows_for(report):
    """Yield (category_name, total_qty, total_value, items) for the full
    detail table. A generator so the item lists are built one category at
    a time and released as soon as they are written."""
    for cat in report.categories:
        if cat.total_qty <= 0:
            continue
        yield (cat.category.name, cat.total_qty, cat.total_value, cat.sorted_items())


def _safe_sheet_name(name, taken):
    safe = (name or "صيدلية")[:31]
    for ch in "[]:*?/\\":
        safe = safe.replace(ch, "-")
    base, n = safe, 1
    while safe in taken:
        n += 1
        safe = f"{base[:28]}-{n}"
    taken.add(safe)
    return safe


def build_workbook(upload, report):
    wb = Workbook(write_only=True)
    has_prices = report.has_prices

    subtitle = (
        f"إجمالي الكمية المتتبعة: {report.total_tracked_qty:g} عبوة   |   "
        f"من الصنف المستهدف: {report.total_target_qty:g} ({report.target_pct:.1f}%)"
    )
    if has_prices:
        subtitle += (
            f"   |   إجمالي القيمة: {report.total_tracked_value:,.0f} {CURRENCY_LABEL}   |   "
            f"أصناف مستهدفة (ل.س.ج): {report.total_target_value:,.0f} ({report.target_value_pct:.1f}%)"
        )
        if not report.has_own_prices:
            subtitle += "   |   (الأسعار موروثة من دفعة شهر سابق)"

    # --- Overall summary sheet ---------------------------------------------
    headers = ["#", "التصنيف", "إجمالي الكمية", "% الكمية", "كمية الصنف المستهدف",
               "نسبة الالتزام % (كمية)"]
    if has_prices:
        headers += [
            f"إجمالي القيمة ({CURRENCY_LABEL})",
            "% القيمة",
            f"قيمة الصنف المستهدف ({CURRENCY_LABEL})",
            "نسبة الالتزام % (قيمة)",
        ]
    summary_widths = [6, 34, 15, 13, 18, 16] + ([18, 12, 20, 16] if has_prices else [])
    sw = SheetWriter(wb, "ملخص التصنيفات", len(headers),
                     widths=summary_widths, freeze=FREEZE_AT)
    _sheet_title(sw, f"ملخص المبيعات حسب التصنيف  |  {upload.name}", subtitle)
    sw.header(headers)
    center = set(range(1, len(headers) + 1)) - {2}
    for i, cat in enumerate(report.categories, start=1):
        if cat.total_qty <= 0:
            continue
        pct_total = (cat.total_qty / report.total_tracked_qty * 100) if report.total_tracked_qty else 0
        values = [i, cat.category.name, round(cat.total_qty, 2), round(pct_total, 1),
                  round(cat.target_qty, 2), round(cat.target_pct, 1)]
        if has_prices:
            pct_value_total = (
                (cat.total_value / report.total_tracked_value * 100)
                if report.total_tracked_value else 0
            )
            values += [
                round(cat.total_value, 0),
                round(pct_value_total, 1),
                round(cat.target_value, 0),
                round(cat.target_value_pct, 1),
            ]
        sw.row(values, fill=(STRIPE_FILL if i % 2 == 0 else NO_FILL), center_cols=center)

    # --- Full detail sheet --------------------------------------------------
    detail_cols = len(_detail_headers(has_prices))
    sw2 = SheetWriter(wb, "الإجمالي العام", detail_cols,
                      widths=_detail_widths(has_prices), freeze=FREEZE_AT)
    _sheet_title(sw2, f"التفصيل الكامل حسب الصنف  |  {upload.name}", subtitle)
    _write_table(sw2, _rows_for(report), has_prices)

    # --- One sheet per pharmacy ---------------------------------------------
    pharmacy_names = set()
    for cat in report.categories:
        pharmacy_names.update(cat.pharmacies.keys())

    taken = set(wb.sheetnames)
    for pharmacy_name in sorted(pharmacy_names):
        rows = []
        ph_total_all = ph_target_all = ph_value_all = ph_target_value_all = 0.0
        for cat in report.categories:
            ph = cat.pharmacies.get(pharmacy_name)
            if ph is None or ph.total_qty == 0:
                continue
            ph_total_all += ph.total_qty
            ph_target_all += ph.target_qty
            ph_value_all += ph.total_value
            ph_target_value_all += ph.target_value
            rows.append((cat.category.name, ph.total_qty, ph.total_value,
                         cat.sorted_items_for_pharmacy(pharmacy_name)))
        ph_pct = (ph_target_all / ph_total_all * 100) if ph_total_all else 0
        ph_sub = (f"إجمالي الكمية: {ph_total_all:g} عبوة   |   "
                  f"من الصنف المستهدف: {ph_target_all:g} ({ph_pct:.1f}%)")
        if has_prices:
            ph_value_pct = (ph_target_value_all / ph_value_all * 100) if ph_value_all else 0
            ph_sub += (f"   |   إجمالي القيمة: {ph_value_all:,.0f} {CURRENCY_LABEL} "
                       f"({ph_value_pct:.1f}% مستهدف)")
        sw_p = SheetWriter(wb, _safe_sheet_name(pharmacy_name, taken), detail_cols,
                           widths=_detail_widths(has_prices), freeze=FREEZE_AT)
        _sheet_title(sw_p, f"{pharmacy_name}  |  {upload.name}", ph_sub)
        _write_table(sw_p, rows, has_prices)
        rows = None

    # --- Missing-price diagnostics sheet -------------------------------------
    if has_prices and report.items_missing_price:
        sw3 = SheetWriter(wb, "أصناف بلا سعر", 2, widths=[45, 16], freeze=FREEZE_AT)
        _sheet_title(
            sw3,
            f"أصناف مُتتبَّعة بلا سعر معروف  |  {upload.name}",
            "قيمتها احتُسبت صفراً في كل الجداول أعلاه حتى يُضاف سعرها",
        )
        sw3.header(["الصنف", "الكمية المباعة"])
        for i, row in enumerate(report.items_missing_price):
            sw3.row(
                [row["item_name"], round(row["quantity"], 2)],
                fill=(WARN_FILL if i % 2 == 0 else NO_FILL),
                center_cols={2},
            )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
