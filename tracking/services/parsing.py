"""Tolerant .xlsx readers for the two workbook shapes this app understands:

1. The "موسع" monthly sales export - one row per sale, with the item name,
   the pharmacy (cost centre), the quantity, and (for some rows) an existing
   category tag (عمود "الشروحات 5").

2. The "تارغيت + كاتيغوري" workbook - a two column lookup of
   item name -> category name, listing the specific products the owner
   instructed pharmacies to sell for each category that month.

Column headers vary a little from month to month (extra spaces, a
trailing "5", Excel merged header rows, etc.), so headers are matched by
keyword rather than by exact string.
"""

from __future__ import annotations

from dataclasses import dataclass

import openpyxl


class WorkbookFormatError(ValueError):
    """Raised when a required column can't be found in the uploaded file."""


def _norm(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _find_column(header_cells: list[str], keywords: list[str]) -> int | None:
    """Return the index of the first header cell containing any keyword."""
    for kw in keywords:
        for idx, cell in enumerate(header_cells):
            if kw in cell:
                return idx
    return None


def _first_sheet_with_header(workbook, required_keyword_sets: list[list[str]], max_scan_rows=8, prefer_name_contains=None):
    """Find a worksheet (and its header row) whose header contains at least
    one keyword from every required set. When several sheets match (e.g. a
    grand-total sheet plus one sheet per pharmacy, all with the same column
    layout), a sheet whose name contains `prefer_name_contains` wins, so we
    don't accidentally sum a single pharmacy's numbers instead of the total."""
    candidates = []
    for ws in workbook.worksheets:
        for row_idx in range(1, max_scan_rows + 1):
            row = [_norm(c) for c in next(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True), [])]
            if not any(row):
                continue
            if all(_find_column(row, kws) is not None for kws in required_keyword_sets):
                candidates.append((ws, row_idx, row))
                break
    if not candidates:
        return None, None, None
    if prefer_name_contains:
        for ws, row_idx, row in candidates:
            if any(kw in ws.title for kw in prefer_name_contains):
                return ws, row_idx, row
    return candidates[0]


def _column_looks_textual(ws, header_row_idx, col_idx, sample_rows=15) -> bool:
    """True when a column's data cells are mostly text (pharmacy names),
    as opposed to numbers (a percentage column, say) - used to tell a
    real "pharmacy" column apart from an unrelated column whose header
    merely contains the word "الصيدلية" (e.g. "% من إجمالي الصيدلية")."""
    seen = 0
    textual = 0
    for row in ws.iter_rows(
        min_row=header_row_idx + 1, max_row=header_row_idx + sample_rows, values_only=True
    ):
        if row is None or col_idx >= len(row):
            continue
        val = row[col_idx]
        if val in (None, ""):
            continue
        seen += 1
        if isinstance(val, str):
            textual += 1
    return seen > 0 and (textual / seen) >= 0.6


def detect_workbook_kind(file_obj) -> str:
    """Sniff whether an uploaded .xlsx is a موسع-style sales file or a
    price/value report (e.g. تقرير أصناف الأورانج), purely from which
    columns its header row has - so the person can select both files
    together in one go without labelling which is which. Returns
    "sales", "prices", or "unknown". Leaves file_obj's read position
    at the start when done.

    Column keywords alone aren't always enough: a price report's
    per-pharmacy sheets have a "% من إجمالي الصيدلية" column, whose
    "الصيدلية" false-matches PHARMACY_KEYWORDS, and some months' موسع
    files add their own "السعر" column, which matches VALUE_KEYWORDS.
    So a sales-shape match is only trusted once its "pharmacy" column's
    actual data looks like pharmacy names (text) rather than numbers -
    that's what a genuine per-sale transaction row looks like and a
    price report's aggregated rows never do."""
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    try:
        sales_ws, sales_header_row, sales_header = _first_sheet_with_header(
            wb, [ITEM_KEYWORDS, PHARMACY_KEYWORDS, QUANTITY_KEYWORDS]
        )
        if sales_ws is not None:
            pharmacy_col = _find_column(sales_header, PHARMACY_KEYWORDS)
            if pharmacy_col is not None and _column_looks_textual(sales_ws, sales_header_row, pharmacy_col):
                return "sales"

        price_ws, _, _ = _first_sheet_with_header(wb, [ITEM_KEYWORDS, QUANTITY_KEYWORDS, VALUE_KEYWORDS])
        if price_ws is not None:
            return "prices"
        return "unknown"
    finally:
        wb.close()
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)


@dataclass
class SalesRecord:
    item_name: str
    pharmacy_name: str
    quantity: float
    raw_category: str


ITEM_KEYWORDS = ["اسم المادة", "اسم الصنف", "الصنف"]
PHARMACY_KEYWORDS = ["مركز الكلفة", "اسم الصيدلية", "الصيدلية"]
QUANTITY_KEYWORDS = ["عبوة", "الكمية", "الكميه"]
CATEGORY_KEYWORDS = ["الشروحات", "التصنيف", "المجموعة", "كاتيغوري", "كاتغوري"]
VALUE_KEYWORDS = ["القيمة", "السعر", "قيمة"]
# A per-row *unit* price column, as opposed to VALUE_KEYWORDS above which
# matches a *total* value column in an aggregated report. Some months'
# موسع files include their own "السعر" column (unit price) right next to
# "عبوة" - see extract_embedded_prices() below.
UNIT_PRICE_KEYWORDS = ["السعر"]


def parse_sales_workbook(file_obj) -> list[SalesRecord]:
    """Parse a موسع-style workbook into a flat list of SalesRecord rows."""
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    ws, header_row_idx, header = _first_sheet_with_header(
        wb, [ITEM_KEYWORDS, PHARMACY_KEYWORDS, QUANTITY_KEYWORDS]
    )
    if ws is None:
        raise WorkbookFormatError(
            "لم أتمكن من العثور على أعمدة (اسم المادة / مركز الكلفة / عبوة) في الملف. "
            "تأكد من أنك ترفع ملف موسع الشهر الصحيح."
        )

    item_col = _find_column(header, ITEM_KEYWORDS)
    pharmacy_col = _find_column(header, PHARMACY_KEYWORDS)
    qty_col = _find_column(header, QUANTITY_KEYWORDS)
    cat_col = _find_column(header, CATEGORY_KEYWORDS)

    records: list[SalesRecord] = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row is None or len(row) <= item_col:
            continue
        item_name = _norm(row[item_col])
        if not item_name:
            continue
        pharmacy = _norm(row[pharmacy_col]) if pharmacy_col is not None and pharmacy_col < len(row) else ""
        raw_qty = row[qty_col] if qty_col is not None and qty_col < len(row) else 0
        try:
            quantity = float(raw_qty) if raw_qty not in (None, "") else 0.0
        except (TypeError, ValueError):
            quantity = 0.0
        raw_category = ""
        if cat_col is not None and cat_col < len(row):
            raw_category = _norm(row[cat_col])
        records.append(
            SalesRecord(
                item_name=item_name,
                pharmacy_name=pharmacy,
                quantity=quantity,
                raw_category=raw_category,
            )
        )
    wb.close()
    return records


def extract_embedded_prices(file_obj) -> list["PriceRecord"]:
    """Some months' موسع files include their own per-row "السعر" (unit
    price) column right next to "عبوة" (quantity) - when that's the case,
    prices for this exact month can be derived straight from the same
    file, with no separate price report needed at all.

    Important: such files often *also* have a "مستهلك" column that looks
    like it should be the row's total value, but it isn't - it stays a
    constant (== السعر × 100) no matter what "عبوة" is for that row (even
    for return rows with negative quantity), so it cannot be used as a
    value. "السعر" itself, however, matches finance's own per-unit price
    almost exactly (cross-checked against أصناف الأورانج report totals),
    so the row's real value is عبوة × السعر, and this function derives a
    weighted-average unit price per item name the same way
    parse_price_report_workbook() does from a separate report.

    The weighting must use the row's *signed* quantity, not its absolute
    value: موسع files routinely include return rows (negative عبوة), and
    weighting by the magnitude of a return inflates the denominator
    without correspondingly reducing the numerator, quietly skewing the
    average price down for any item with returns - which then throws off
    every value figure computed from it later (reporting.py multiplies
    each row's own signed quantity by this single average price, so the
    average has to be the one signed-weighted average that reproduces the
    item's true net value when multiplied by its true net quantity).

    Returns an empty list when the file has no such column (e.g. an older
    موسع export, or a month where the column is simply absent)."""
    if hasattr(file_obj, "seek"):
        file_obj.seek(0)
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    try:
        ws, header_row_idx, header = _first_sheet_with_header(
            wb, [ITEM_KEYWORDS, PHARMACY_KEYWORDS, QUANTITY_KEYWORDS]
        )
        if ws is None:
            return []

        item_col = _find_column(header, ITEM_KEYWORDS)
        qty_col = _find_column(header, QUANTITY_KEYWORDS)
        price_col = _find_column(header, UNIT_PRICE_KEYWORDS)
        if price_col is None:
            return []

        totals: dict[str, list[float]] = {}
        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if row is None or len(row) <= max(item_col, qty_col, price_col):
                continue
            item_name = _norm(row[item_col])
            if not item_name:
                continue
            try:
                qty = float(row[qty_col])
                price = float(row[price_col])
            except (TypeError, ValueError):
                continue
            if qty == 0 or price <= 0:
                continue
            entry = totals.setdefault(item_name, [0.0, 0.0])
            entry[0] += qty * price
            entry[1] += qty
        return [
            PriceRecord(item_name=name, unit_price=value_sum / qty_sum)
            for name, (value_sum, qty_sum) in totals.items()
            if qty_sum != 0
        ]
    finally:
        wb.close()
        if hasattr(file_obj, "seek"):
            file_obj.seek(0)


@dataclass
class TargetRecord:
    item_name: str
    category_name: str


def parse_target_workbook(file_obj) -> list[TargetRecord]:
    """Parse a تارغيت+كاتيغوري style workbook: item -> category lookup."""
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    ws, header_row_idx, header = _first_sheet_with_header(wb, [ITEM_KEYWORDS, CATEGORY_KEYWORDS])
    if ws is None:
        raise WorkbookFormatError(
            "لم أتمكن من العثور على عمودي (الصنف / التصنيف) في الملف. "
            "تأكد من أنك تختار ملف التارغيت كاتيغوري الصحيح."
        )

    item_col = _find_column(header, ITEM_KEYWORDS)
    cat_col = _find_column(header, CATEGORY_KEYWORDS)

    records: list[TargetRecord] = []
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row is None or len(row) <= max(item_col, cat_col):
            continue
        item_name = _norm(row[item_col])
        category_name = _norm(row[cat_col])
        if not item_name or not category_name:
            continue
        records.append(TargetRecord(item_name=item_name, category_name=category_name))
    wb.close()
    return records


@dataclass
class PriceRecord:
    item_name: str
    unit_price: float


def parse_price_report_workbook(file_obj) -> list[PriceRecord]:
    """Parse a detailed report workbook (e.g. تقرير أصناف الأورانج) that
    lists, per item, a quantity column and a total-value column (القيمة
    (ل.س)). The unit price is derived as value / quantity, weighted across
    every matching row for that item name (so it still works even if an
    item appears more than once). Prefers a "grand total" sheet (its name
    containing "اجمالي"/"إجمالي") over a single pharmacy's sheet, so
    quantities aren't double counted or scoped to one branch."""
    wb = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    ws, header_row_idx, header = _first_sheet_with_header(
        wb, [ITEM_KEYWORDS, QUANTITY_KEYWORDS, VALUE_KEYWORDS], prefer_name_contains=["اجمالي", "إجمالي"]
    )
    if ws is None:
        raise WorkbookFormatError(
            "لم أتمكن من العثور على أعمدة (الصنف / الكمية / القيمة) في الملف. "
            "تأكد من أنك ترفع تقريراً يحتوي على قيمة لكل صنف."
        )

    item_col = _find_column(header, ITEM_KEYWORDS)
    qty_col = _find_column(header, QUANTITY_KEYWORDS)
    value_col = _find_column(header, VALUE_KEYWORDS)

    totals: dict[str, list[float, float]] = {}
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row is None or len(row) <= max(item_col, qty_col, value_col):
            continue
        item_name = _norm(row[item_col])
        if not item_name:
            continue
        try:
            qty = float(row[qty_col])
            value = float(row[value_col])
        except (TypeError, ValueError):
            continue
        if qty <= 0:
            continue
        entry = totals.setdefault(item_name, [0.0, 0.0])
        entry[0] += value
        entry[1] += qty
    wb.close()

    return [
        PriceRecord(item_name=name, unit_price=value_sum / qty_sum)
        for name, (value_sum, qty_sum) in totals.items()
        if qty_sum > 0
    ]
